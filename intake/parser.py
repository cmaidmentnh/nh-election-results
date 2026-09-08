"""Claude-backed extraction of vote totals from free-form reports.

Two calls per message:

  1. identify_town - which polling place is this? (skipped when the town name
     matched exactly in Python, which is the common case)
  2. extract      - given that town's exact ballot, attach numbers to race ids
                    and candidate ids

The model is never asked to invent a race or a candidate. It picks from the
ids it is given, and anything it cannot place comes back unmatched so a human
sees it rather than the parser guessing.
"""

import base64
import logging
import re
import mimetypes
from pathlib import Path

import anthropic
from pydantic import BaseModel, Field

from intake import config

log = logging.getLogger("intake.parser")

_client = None


def client():
    global _client
    if _client is None:
        _client = anthropic.Anthropic(api_key=config.ANTHROPIC_API_KEY or None)
    return _client


class TownGuess(BaseModel):
    municipality: str = Field(description="Exact polling place name from the list, or empty string if unclear")
    confidence: float = Field(description="0.0 to 1.0")
    reasoning: str


class VoteLine(BaseModel):
    race_id: int = Field(description="race_id from the ballot, or 0 if none fits")
    candidate_id: int = Field(description="candidate_id from the ballot, or 0 if none fits")
    candidate_text: str = Field(description="The name exactly as the reporter wrote it")
    race_text: str = Field(description="The race exactly as the reporter wrote it, if stated")
    votes: int
    confidence: float = Field(description="0.0 to 1.0 for this one line")


class BallotsLine(BaseModel):
    election_id: int = Field(description="election_id whose ballots these are, or 0 if unclear")
    ballots_cast: int
    confidence: float


class RaceCheck(BaseModel):
    race_id: int = Field(description="race_id this checksum belongs to")
    stated_total: int = Field(description="The TOTAL VOTES CAST printed for this race, or 0 if the sheet does not print one")
    blanks: int = Field(description="BLANKS or UNDERVOTES printed for this race, or 0")
    overvotes: int = Field(description="OVERVOTES printed for this race, or 0")


class Extraction(BaseModel):
    municipality_text: str = Field(description="The place name as the reporter wrote it")
    contains_results: bool = Field(description="False for chatter, questions, acknowledgements")
    notes: str = Field(description="Anything a human reviewer should know; empty if nothing")
    report_scope: str = Field(
        description="One of: 'machine_tape' (one machine's tape, part of the town), "
                    "'town_total' (the town's combined total for these races), "
                    "'unknown'. Judge from the sheet itself; do not guess.")
    device_label: str = Field(
        description="Any machine/device/tape identifier printed on the report "
                    "(e.g. 'Machine 1 of 2', 'AccuVote #3', 'District 2 tape'), "
                    "empty string if none. Never invent one.")
    lines: list[VoteLine]
    ballots: list[BallotsLine]
    checks: list[RaceCheck] = Field(
        description="One entry per race whose sheet prints a total, for arithmetic verification")


TOWN_SYSTEM = """You identify which New Hampshire polling place an election-night \
report is about.

Reply with the polling place name copied EXACTLY from the list you are given, or \
an empty string if you cannot tell. Never invent a name or return one that is not \
on the list.

New Hampshire towns are often confused with each other. Take care with:
- Ward-split cities (Manchester, Nashua, Concord, Dover, Rochester, Keene, Laconia, \
Portsmouth, Somersworth, Berlin, Claremont, Franklin, Lebanon). A report from one of \
these must name the ward. Without a ward number, return an empty string.
- Similar names: Hampton / Hampton Falls / North Hampton / East Kingston / Kingston, \
Newport / Newington / Newmarket / New Ipswich / Newton, Lyme / Lyman / Lyndeborough."""

EXTRACT_SYSTEM = """You read election-night vote reports from New Hampshire town \
clerks and party volunteers and turn them into structured vote totals.

You are given the EXACT ballot for one polling place: every race with its race_id, \
and every candidate with its candidate_id. Your only job is to attach the reported \
numbers to those ids.

THIS IS A PRIMARY. THERE ARE TWO SEPARATE BALLOTS.
The Republican ballot and the Democratic ballot are different elections held the \
same day. The same office appears on BOTH, with a DIFFERENT race_id and a \
different set of candidates. Governor on the Republican ballot and Governor on the \
Democratic ballot are not the same race. Choosing the wrong one files a \
candidate's votes under the wrong party, which is worse than not filing them at \
all. So:
- A report usually states its party once, at the top ("Republican primary", "GOP", \
  "R side", "Dem results"). That party applies to every line in the report until \
  the report says otherwise.
- If the party is never stated, use the ballot the reported candidates actually \
  appear on - the two candidate lists do not overlap.
- If a line could sit on either ballot and nothing resolves it, return race_id=0 \
  and let a human decide. Never guess between the two.

Every race listed in the ballot above IS on this town's ballot. If the report names \
an office you can see above - including abbreviations like "Exec Council D4", "EC4", \
"CD1", "Sheriff", "Reg of Deeds" - use that race_id. Do not claim a race is absent \
when it is listed.

Other rules:
- Only ever use a race_id and candidate_id that appear in the ballot above. Never \
invent one.
- If a reported name does not match any candidate on that ballot, still return the \
line with candidate_id=0 and the name in candidate_text. A human will resolve it. \
Do NOT force it onto the nearest name.
- Match on surname where the ballot is unambiguous; if two candidates in the same \
race share a surname, only match when the report distinguishes them, otherwise \
return candidate_id=0.
- Votes written as "412" and "412 votes" and "412 (52%)" are all 412. Never derive a \
vote count from a percentage.
- A total ballots cast / turnout figure is NOT a candidate vote. Put it in ballots, \
using the election_id shown in that ballot's header above. A primary reports \
ballots cast per party, so "1,420 Republican ballots" uses the Republican \
election_id.
- For every race where the sheet prints a TOTAL VOTES CAST (and/or BLANKS,
UNDERVOTES, OVERVOTES), copy those figures into checks for that race_id. Copy
them exactly as printed; do not compute them yourself and do not infer a total
that is not on the page. They are used to verify the candidate numbers add up,
so a guessed total defeats the check. Leave a figure as 0 when the sheet does
not print it.
- Set confidence per line. Use a low value when handwriting is unclear, when a digit \
is ambiguous, or when the name match was a stretch.
- If the message contains no vote totals at all (a question, a greeting, "on my way"), \
set contains_results=false and return no lines.
REPORTS ARRIVE IN TWO SHAPES, AND WHICH ONE MATTERS.
A town may send each machine's tape first and a combined town total afterwards.
Set report_scope to 'machine_tape' when the sheet is one device's output (it
names a machine, device or tape number, or is plainly a printer tape), and
'town_total' when it is the town's combined return for those races (a Return of
Votes form, a sheet headed with the town name and totals, or one that reconciles
to the town's ballots cast). Use 'unknown' when the sheet does not make it clear
- do not guess, because the two are combined differently downstream.

MANY TOWNS COUNT ON SEVERAL MACHINES AND SEND ONE TAPE PER MACHINE.
A tape may therefore hold only part of a town's vote. If the report names a
machine, device, or tape number, put it in device_label exactly as printed.
Never add tapes together yourself and never scale a number up to what you think
the town total should be - report only what this sheet shows.

PHOTOGRAPHS AND SCANNED TALLY SHEETS ARE THE NORMAL CASE.
Most reports are a phone photo of the machine tape or the hand-tally sheet, or a
scanned PDF of it, with little or no typed text. Read them carefully:
- Work down the sheet race by race. A New Hampshire tally sheet lists the office, \
then each candidate with a vote total beside or beneath the name.
- The sheet usually says which ballot it is - "REPUBLICAN", "DEMOCRATIC", "REP", \
"DEM", often in the header or as a column heading. Use it. If one sheet has \
separate Republican and Democratic columns, read each into its own ballot's \
race_id.
- These lines are NOT candidates and must never be returned as candidate votes: \
TOTAL, TOTAL VOTES CAST, BLANKS, BLANK, UNDERVOTES, OVERVOTES, SCATTERING, \
SCATTERED, VOID, SPOILED, ABSENTEE (as a column heading), REGISTERED VOTERS. \
A "TOTAL BALLOTS CAST" or "BALLOTS CAST" figure goes in ballots, not lines.
- "SCATTERING" or a plain "WRITE-IN" total is the aggregate write-in line - use \
the "Write-in (aggregate line)" candidate_id for that race.
- Multi-seat State Representative races list many candidates at once; report every \
one you can read.
- Numbers are often handwritten. Distinguish carefully between 1/7, 3/8, 5/6, and \
0/6/8. If a digit is genuinely ambiguous, lower the confidence for that line - do \
not silently pick one.
- Read only what is printed. If a number is cut off, smudged, obscured by glare, or \
outside the frame, omit that line and say so in notes rather than guessing.
- If the photograph is too blurry or too dark to read at all, set \
contains_results=false and say so in notes."""


# Claude Opus 5 reads up to 2576px on the long edge. Tally tapes are dense
# columns of small digits, so send at that ceiling rather than downscaling
# further - and no larger, which would only cost tokens.
MAX_EDGE = 2576
MAX_ATTACHMENTS = 12
# A long tape becomes several tiles; cap the total so one photo cannot
# swamp a request.
MAX_TILES = 6
MAX_PDF_BYTES = 20 * 1024 * 1024


def _sniff(data):
    """Content type from magic bytes.

    Filenames cannot be trusted here: signal-cli stores attachments under a
    bare id with no extension, so guessing by suffix skips every photo posted
    to the Signal group.
    """
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    if data[:5] == b"%PDF-":
        return "application/pdf"
    if data[4:8] == b"ftyp" and data[8:12] in (b"heic", b"heix", b"hevc", b"mif1", b"msf1"):
        return "image/heic"
    if data[:4] == b"PK\x03\x04":
        # A zip container. .xlsx is the one that matters: at least one town sent
        # its results as a spreadsheet in 2024.
        return "application/zip"
    return None


def _prepare_image(data, media_type):
    """Re-encode to JPEG within the model's resolution ceiling.

    Returns a LIST of images: a long tally tape shrunk to fit 2576px on its
    long edge would have its digits squashed into illegibility, so tall images
    are cut into overlapping tiles at full readable width instead. The overlap
    keeps a row that falls on a seam readable in one tile or the other.

    Phone photos arrive at 12MP and, from iPhones, often as HEIC, which the API
    does not accept; both are handled here so the caller does not care.
    """
    try:
        import io

        from PIL import Image
        try:  # iPhone photos
            import pillow_heif
            pillow_heif.register_heif_opener()
        except ImportError:
            if media_type == "image/heic":
                return []

        img = Image.open(io.BytesIO(data))
        img.load()
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")

        # Scale by WIDTH so text keeps its size, then tile down the length.
        if img.width > MAX_EDGE:
            ratio = MAX_EDGE / img.width
            img = img.resize((MAX_EDGE, max(1, int(img.height * ratio))), Image.LANCZOS)

        pieces = []
        if img.height <= MAX_EDGE:
            pieces.append(img)
        else:
            overlap = 220
            step = MAX_EDGE - overlap
            top = 0
            while top < img.height and len(pieces) < MAX_TILES:
                bottom = min(top + MAX_EDGE, img.height)
                pieces.append(img.crop((0, top, img.width, bottom)))
                if bottom >= img.height:
                    break
                top += step

        out = []
        for piece in pieces:
            buf = io.BytesIO()
            piece.save(buf, format="JPEG", quality=88, optimize=True)
            out.append((buf.getvalue(), "image/jpeg"))
        return out
    except Exception:
        log.exception("Could not re-encode an attachment")
        if media_type in ("image/jpeg", "image/png", "image/gif", "image/webp"):
            return [(data, media_type)]
        return []


def _spreadsheet_text(path):
    """An .xlsx of results rendered as plain rows the model can read."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        log.warning("openpyxl not installed; cannot read %s", path)
        return ""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        log.exception("Could not open spreadsheet %s", path)
        return ""
    out = []
    for ws in wb.worksheets:
        out.append(f"--- sheet: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                out.append(" | ".join(cells).rstrip(" |"))
            if len(out) > 4000:
                out.append("(truncated)")
                break
    wb.close()
    return "\n".join(out)


def spreadsheet_texts(paths):
    """Text extracted from any spreadsheet attachments, for the prompt body."""
    chunks = []
    for p in (paths or []):
        path = Path(p)
        try:
            if not path.exists():
                continue
            head = path.open("rb").read(8)
        except OSError:
            continue
        if head[:4] != b"PK\x03\x04" and path.suffix.lower() not in (".xlsx", ".xlsm"):
            continue
        text = _spreadsheet_text(path)
        if text:
            chunks.append(f"[spreadsheet attachment: {path.name}]\n{text}")
    return "\n\n".join(chunks)


def _attachment_blocks(paths, limit=MAX_ATTACHMENTS):
    """Photos and PDFs of tally sheets as content blocks.

    Most reports on election night are a picture of the tape or a scanned
    printout rather than typed numbers, so this is the main input path.
    """
    blocks = []
    for p in (paths or [])[:limit]:
        path = Path(p)
        try:
            if not path.exists() or path.stat().st_size == 0:
                continue
            data = path.read_bytes()
        except OSError:
            log.exception("Could not read attachment %s", p)
            continue

        media_type = _sniff(data) or mimetypes.guess_type(str(path))[0]

        if media_type == "application/pdf":
            if len(data) > MAX_PDF_BYTES:
                log.warning("Skipping oversized PDF %s (%.1f MB)", p, len(data) / 1e6)
                continue
            blocks.append({
                "type": "document",
                "source": {"type": "base64", "media_type": "application/pdf",
                           "data": base64.standard_b64encode(data).decode()},
            })
            continue

        if not media_type or not media_type.startswith("image/"):
            continue

        for prepared, out_type in _prepare_image(data, media_type):
            blocks.append({
                "type": "image",
                "source": {"type": "base64", "media_type": out_type,
                           "data": base64.standard_b64encode(prepared).decode()},
            })
    return blocks


def identify_town(town_list, body, subject="", sender="", attachments=None):
    content = _attachment_blocks(attachments)
    content.append({
        "type": "text",
        "text": (
            f"Polling places:\n{town_list}\n\n"
            f"--- REPORT ---\nFrom: {sender}\nSubject: {subject}\n\n{body or '(no text; see image)'}"
        ),
    })
    resp = client().messages.parse(
        model=config.MODEL,
        max_tokens=4000,
        system=TOWN_SYSTEM,
        messages=[{"role": "user", "content": content}],
        output_format=TownGuess,
    )
    return resp.parsed_output


def extract(municipality, roster_text, body, subject="", sender="", attachments=None,
            max_tokens=None, _retry=True):
    content = _attachment_blocks(attachments)
    content.append({
        "type": "text",
        "text": (
            f"POLLING PLACE: {municipality}\n"
            f"BALLOT FOR THIS POLLING PLACE:\n{roster_text}\n\n"
            f"--- REPORT ---\nFrom: {sender}\nSubject: {subject}\n\n"
            f"{body or '(no text; results are in the attached file)'}"
        ),
    })
    sheets = spreadsheet_texts(attachments)
    if sheets:
        content.append({"type": "text", "text": f"\n--- ATTACHED SPREADSHEET ---\n{sheets}"})
    try:
        resp = client().messages.parse(
            model=config.MODEL,
            max_tokens=max_tokens or config.MAX_OUTPUT_TOKENS,
            system=EXTRACT_SYSTEM,
            thinking={"type": "adaptive"},
            messages=[{"role": "user", "content": content}],
            output_format=Extraction,
        )
        return resp.parsed_output
    except Exception as exc:
        # A big return-of-votes sheet can run past the output ceiling and come
        # back as truncated JSON. Losing the whole town over that is the worst
        # outcome, so retry once with far more room - which the SDK only allows
        # when streaming.
        if _retry and "json" in str(exc).lower():
            log.warning("%s: reply did not parse (%s); retrying streamed",
                        municipality, str(exc)[:120])
            return _extract_streamed(content)
        raise


def _extract_streamed(content):
    """One extraction with a large output ceiling, which requires streaming."""
    with client().messages.stream(
        model=config.MODEL,
        max_tokens=config.MAX_OUTPUT_TOKENS_RETRY,
        system=EXTRACT_SYSTEM,
        thinking={"type": "adaptive"},
        messages=[{"role": "user", "content": content}],
        output_config={"format": {"type": "json_schema",
                                  "schema": Extraction.model_json_schema()}},
    ) as stream:
        message = stream.get_final_message()
    text = next(b.text for b in message.content if b.type == "text")
    return Extraction.model_validate_json(text)


def line_key(line):
    """Identity of one reported line across independent reads.

    Matched lines are identified by the ids they resolved to. Unmatched lines
    have race_id=0 and candidate_id=0, so keying on ids alone would collapse
    every unmatched line in a report into a single bucket - and then a value
    agreed for one of them would be written onto all of them. Fall back to the
    text the reporter actually wrote, which is what distinguishes them.
    """
    if line.race_id and line.candidate_id:
        return ("id", line.race_id, line.candidate_id)
    # Normalise hard: two reads of the same page often differ only in
    # punctuation ("Donald J. Trump" vs "Donald J Trump"), and treating those
    # as different lines would report a disagreement that is not one.
    def norm(t):
        return re.sub(r"[^a-z0-9]+", "", (t or "").lower())

    return ("text", norm(line.race_text), norm(line.candidate_text))


def _tally(reads):
    """Per (race, candidate), what each read said. Missing counts against it."""
    votes = {}
    for r in reads:
        seen_here = {}
        for l in r.lines:
            seen_here.setdefault(line_key(l), l.votes)
        for key, v in seen_here.items():
            votes.setdefault(key, [])
    for key in votes:
        for r in reads:
            match = next((l.votes for l in r.lines if line_key(l) == key), None)
            votes[key].append(match)
    return votes


def _verdict(values):
    """(winning_value, agreed?) for one line's readings across every pass.

    A supermajority is required, never a bare majority. Repeated reading fixes
    RANDOM misreads - a smudged digit read differently each time - but it does
    not fix SYSTEMATIC ones. If the model consistently reads the wrong column of
    a layout it will do so every pass, and a bare 3-2 majority would then
    manufacture false confidence in a wrong number. So a close split is treated
    as unresolved and sent to a human, which is the honest answer.
    """
    present = [v for v in values if v is not None]
    if not present:
        return None, False
    counts = {}
    for v in present:
        counts[v] = counts.get(v, 0) + 1
    ranked = sorted(counts.items(), key=lambda kv: -kv[1])
    top_value, top_count = ranked[0]
    runner_up = ranked[1][1] if len(ranked) > 1 else 0
    n = len(values)

    if n <= 2:
        return top_value, top_count == n            # both reads, same answer
    if top_count >= 4:
        return top_value, True                      # 4+ of 5 agree
    if top_count >= 3 and runner_up <= 1:
        return top_value, True                      # 3-1-1, one clear answer
    return top_value, False                         # 3-2 or 2-2-1: genuinely split


def extract_consensus(municipality, roster_text, body, subject="", sender="",
                      attachments=None, max_reads=None):
    """Read the report until the readings settle, then report what is solid.

    Back-testing against real 2024 clerk PDFs showed the model reporting high
    confidence on numbers that were badly wrong - Hollis County Sheriff read as
    2,571 against an actual 5,337. Self-reported confidence therefore cannot be
    the only gate.

    Two reads are taken first. If every line agrees, that is the answer. If any
    line disagrees the whole report is read again, up to max_reads, and each
    line is decided by supermajority - see _verdict for why a bare majority is
    not enough. The winning value replaces whatever the first read said, so a
    line that lost 1-4 publishes the number the other four saw.

    Returns (extraction, agreed_keys, disagreements) where disagreements maps a
    key to the full list of readings, for the review queue to show.
    """
    from concurrent.futures import ThreadPoolExecutor

    def read_once():
        return extract(municipality, roster_text, body, subject, sender, attachments)

    # The two baseline reads are independent, so pay for one round trip, not
    # two. On election night latency per report is what decides whether the
    # board keeps up with the towns.
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(read_once) for _ in range(2)]
        reads = [f.result() for f in futures]

    first = reads[0]
    if not first.contains_results or not first.lines:
        return first, set(), {}

    def unsettled(rs):
        return any(not _verdict(v)[1] for v in _tally(rs).values())

    while len(reads) < (max_reads or config.MAX_READS) and unsettled(reads):
        reads.append(extract(municipality, roster_text, body, subject, sender, attachments))

    tally = _tally(reads)
    agreed, disagreements = set(), {}
    for key, values in tally.items():
        winner, ok = _verdict(values)
        if ok:
            agreed.add(key)
        else:
            disagreements[key] = values

        # Publish what the reads actually settled on, not the first attempt.
        if winner is not None:
            for line in first.lines:
                if line_key(line) == key:
                    line.votes = winner

    # A line only later reads saw at all still belongs in the report.
    seen = {line_key(l) for l in first.lines}
    for r in reads[1:]:
        for l in r.lines:
            key = line_key(l)
            if key not in seen and key in tally:
                winner, ok = _verdict(tally[key])
                if winner is not None:
                    l.votes = winner
                    first.lines.append(l)
                    seen.add(key)

    log.info("%s: %d reads, %d settled, %d split",
             municipality, len(reads), len(agreed), len(disagreements))
    return first, agreed, disagreements
